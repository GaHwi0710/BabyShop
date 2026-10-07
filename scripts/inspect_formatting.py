import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import sys

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document('baocaobtl.docx')

print("=== 1. SECTIONS AND MARGINS ===")
for i, s in enumerate(doc.sections):
    print(f"Section {i}: Page size ({s.page_width.inches:.2f}\" x {s.page_height.inches:.2f}\"), "
          f"Margins: Top={s.top_margin.inches:.2f}\", Bottom={s.bottom_margin.inches:.2f}\", "
          f"Left={s.left_margin.inches:.2f}\", Right={s.right_margin.inches:.2f}\"")

print("\n=== 2. PARAGRAPH STYLES AND FONTS ===")
styles_used = set(p.style.name for p in doc.paragraphs)
print("Styles used in paragraphs:", sorted(styles_used))

font_names = set()
font_sizes = set()
for p in doc.paragraphs[:100]:
    for r in p.runs:
        if r.font.name:
            font_names.add(r.font.name)
        if r.font.size:
            font_sizes.add(r.font.size.pt)
print("Font names used:", font_names)
print("Font sizes used:", sorted(font_sizes))

print("\n=== 3. HEADINGS KEEP_WITH_NEXT STATUS ===")
orphan_headings = 0
for i, p in enumerate(doc.paragraphs):
    if p.style.name.startswith("Heading"):
        if not p.paragraph_format.keep_with_next:
            orphan_headings += 1
print(f"Headings without keep_with_next: {orphan_headings}")

print("\n=== 4. EMPTY PARAGRAPHS ===")
empty_p = [i for i, p in enumerate(doc.paragraphs) if not p.text.strip() and len(p._element.xpath('.//w:drawing')) == 0]
print(f"Total completely empty paragraphs: {len(empty_p)}")
print("Sample empty paragraphs indices:", empty_p[:20])

print("\n=== 5. TABLE FORMATTING CHECK ===")
tbl_without_header = 0
for idx, tbl in enumerate(doc.tables):
    # check if tblHeader is set on row 0
    trPr = tbl.rows[0]._tr.get_or_add_trPr()
    has_tbl_header = len(trPr.xpath('.//w:tblHeader')) > 0
    if not has_tbl_header:
        tbl_without_header += 1
print(f"Tables without tblHeader on row 0: {tbl_without_header} / {len(doc.tables)}")
