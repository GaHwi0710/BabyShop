# -*- coding: utf-8 -*-
import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document('D:/BabyShop/Nhom1_BCBTL.docx')

print("=== 1. TỔNG QUAN FILE ===")
print("Paragraphs:", len(doc.paragraphs), "Tables:", len(doc.tables))
print("Trang bìa P4:", doc.paragraphs[4].text)

print("\n=== 2. KIỂM TRA FONT SIZE CÁC TỪ CODE (COURIER NEW) ===")
courier_p_size10 = 0
courier_p_other = 0
for p in doc.paragraphs:
    for r in p.runs:
        if r.font.name == 'Courier New':
            if r.font.size and r.font.size.pt == 10.0:
                courier_p_size10 += 1
            else:
                courier_p_other += 1

courier_t_size10 = 0
courier_t_other = 0
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    if r.font.name == 'Courier New':
                        if r.font.size and r.font.size.pt == 10.0:
                            courier_t_size10 += 1
                        else:
                            courier_t_other += 1

print(f"Courier New runs in Paragraphs: size 10 = {courier_p_size10}, other = {courier_p_other}")
print(f"Courier New runs in Tables: size 10 = {courier_t_size10}, other = {courier_t_other}")

print("\n=== 3. KIỂM TRA MẪU CÁC BẢNG SCHEMA CSDL ===")
t9 = doc.tables[9] # Users
for ri in range(1, 4):
    r = t9.rows[ri]
    r0 = r.cells[0].paragraphs[0].runs[0]
    r1 = r.cells[1].paragraphs[0].runs[0]
    print(f"T9 R{ri}: col0='{r0.text}' (font={r0.font.name}, size={r0.font.size.pt if r0.font.size else None}, bold={r0.font.bold}) | col1='{r1.text}' (font={r1.font.name}, size={r1.font.size.pt if r1.font.size else None})")

t15 = doc.tables[15] # SanPham
for ri in range(len(t15.rows)-2, len(t15.rows)):
    r = t15.rows[ri]
    r0 = r.cells[0].paragraphs[0].runs[0]
    r1 = r.cells[1].paragraphs[0].runs[0]
    print(f"T15 R{ri}: col0='{r0.text}' (font={r0.font.name}, size={r0.font.size.pt if r0.font.size else None}, bold={r0.font.bold}) | col1='{r1.text}' (font={r1.font.name}, size={r1.font.size.pt if r1.font.size else None})")

print("\n=== 4. KIỂM TRA MẪU ĐOẠN VĂN (MỤC 4.6 & 4.7) ===")
for pi in [586, 587, 598]:
    p = doc.paragraphs[pi]
    print(f"P{pi}: {[r.font.size.pt if r.font.size else None for r in p.runs[:4]]} | text='{p.text[:60]}'")

print("\n=== 5. CHẠY LẠI TOÀN BỘ 16 TIÊU CHÍ AUDIT ===")
checks = {}
checks['1_no_reminder'] = not any('Nên thống nhất một kiểu duy nhất' in p.text for p in doc.paragraphs)
checks['2_tables_count_45'] = len(doc.tables) == 45
checks['3_metadata_table'] = any('Bảng 4.32.' in p.text for p in doc.paragraphs)
p_529 = [p for p in doc.paragraphs if p.text.startswith('5.2.9.')]
p_5219 = [p for p in doc.paragraphs if p.text.startswith('5.2.19.')]
checks['4_heading_3'] = (len(p_529) > 0 and p_529[0].style.name == 'Heading 3' and len(p_5219) > 0 and p_5219[0].style.name == 'Heading 3')
checks['5_17_reports'] = any('17. Báo cáo Hàng bán bị đổi trả' in p.text for p in doc.paragraphs)
checks['6_accounting_45_complete'] = any('theo Chế độ kế toán Doanh nghiệp nhỏ và vừa' in p.text for p in doc.paragraphs if 'Sổ nhật ký chung' in p.text)
checks['7_legal_references'] = any('Luật Kế toán số 88/2015/QH13' in p.text for p in doc.paragraphs)

erd_size = 0
for i, p in enumerate(doc.paragraphs):
    if 'Hình 4.1. Sơ đồ quan hệ logic' in p.text:
        prev_p = doc.paragraphs[i-1]
        for r in prev_p.runs:
            for blip in r._r.xpath('.//a:blip'):
                r_id = blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                erd_size = len(doc.part.related_parts[r_id]._blob)
checks['8_erd_image_updated'] = (erd_size == 183592)
checks['9_schedule_12_weeks'] = any('12 tuần của học kỳ' in p.text for p in doc.paragraphs)
checks['10_schedule_table_rows'] = (len(doc.tables[43].rows) == 13)
sp_fields = [r.cells[0].text.strip() for r in doc.tables[15].rows]
checks['11_sp_mavach'] = ('MaVach' in sp_fields and 'BienThe' in sp_fields)
cthd_fields = [r.cells[0].text.strip() for r in doc.tables[21].rows]
checks['12_cthd_giavon'] = ('GiaVon' in cthd_fields and 'ThueSuat' in cthd_fields)
pn_fields = [r.cells[0].text.strip() for r in doc.tables[22].rows]
checks['13_pn_sohdncc'] = ('SoHoaDonNCC' in pn_fields and 'DaThanhToan' in pn_fields)
cn_fields = [r.cells[0].text.strip() for r in doc.tables[35].rows]
checks['14_cn_ngayphatsinh'] = ('NgayPhatSinh' in cn_fields and 'HanThanhToan' in cn_fields)
vt_desc = [r.cells[3].text for r in doc.tables[10].rows if r.cells[0].text.strip() == 'MaVaiTro']
checks['15_vaitro_admin'] = len(vt_desc) > 0 and 'Admin' in vt_desc[0]
checks['16_appendix_script'] = any('ensure-indexes.js' in p.text for p in doc.paragraphs)

for k, v in checks.items():
    print(f"  [{'PASS' if v else 'FAIL'}] {k}")
passed = sum(1 for v in checks.values() if v)
print(f"==> TỔNG KẾT AUDIT: {passed}/{len(checks)} TIÊU CHÍ ĐẠT!")
