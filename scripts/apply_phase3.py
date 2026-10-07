import sys
import os
import docx
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl

sys.stdout.reconfigure(encoding='utf-8')

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    for shd in tcPr.xpath('.//w:shd'):
        tcPr.remove(shd)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    for mar in tcPr.xpath('.//w:tcMar'):
        tcPr.remove(mar)
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def main():
    print("===============================================================")
    print("BẮT ĐẦU PHA 3: HOÀN THIỆN, CHUẨN HÓA VÀ NGHIỆM THU BÁO CÁO")
    print("===============================================================")
    
    src_file = "baocaobtl.docx"
    dst_docx = "baocaobtl_final.docx"
    dst_pdf = "baocaobtl_final.pdf"
    
    doc = docx.Document(src_file)
    print(f"1. Đã nạp tài liệu nguồn {src_file} ({len(doc.paragraphs)} đoạn, {len(doc.tables)} bảng)")

    # =========================================================================
    # BƯỚC 1: XÓA SẠCH "BABYBON" VÀ SỬA CÁC LỖI THUẬT NGỮ CÒN SÓT
    # =========================================================================
    print("\n--- BƯỚC 1: Xóa BabyBon & Chuẩn hóa thuật ngữ ---")
    babybon_fixes = [
        ("Cửa hàng Mẹ và Bé BabyBon", "Cửa hàng Mẹ và Bé"),
        ("BabyBon – Cửa hàng Mẹ & Bé ra đời", "Cửa hàng Mẹ và Bé ra đời"),
        ("BabyBon – Cửa hàng Mẹ & Bé được tổ chức", "Cửa hàng Mẹ và Bé được tổ chức"),
        ("BabyBon – Cửa hàng Mẹ & Bé", "Cửa hàng Mẹ và Bé"),
        ("BabyBon được định hình", "Cửa hàng Mẹ và Bé được định hình"),
        ("BabyBon đang từng bước", "Cửa hàng Mẹ và Bé đang từng bước"),
        ("BabyBon hoạt động chủ yếu", "Cửa hàng Mẹ và Bé hoạt động chủ yếu"),
        ("BabyBon vẫn sở hữu", "Cửa hàng Mẹ và Bé vẫn sở hữu"),
        ("BabyBon chủ yếu áp dụng", "Cửa hàng Mẹ và Bé chủ yếu áp dụng"),
        ("sự tồn tại và phát triển của BabyBon", "sự tồn tại và phát triển của cửa hàng"),
        ("Hệ thống quản lý Cửa hàng Mẹ và Bé BabyBon", "Hệ thống quản lý Cửa hàng Mẹ và Bé"),
        ("tại BabyBon", "tại Cửa hàng Mẹ và Bé"),
        ("của BabyBon", "của Cửa hàng Mẹ và Bé"),
        ("BabyBon", "Cửa hàng Mẹ và Bé"),
    ]
    
    bb_count = 0
    for p in doc.paragraphs:
        orig = p.text
        for old_txt, new_txt in babybon_fixes:
            if old_txt in p.text:
                p.text = p.text.replace(old_txt, new_txt)
                for r in p.runs:
                    r.font.name = "Times New Roman"
        if orig != p.text:
            bb_count += 1

    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for old_txt, new_txt in babybon_fixes:
                    if old_txt in cell.text:
                        cell.text = cell.text.replace(old_txt, new_txt)
                        for p in cell.paragraphs:
                            for r in p.runs:
                                r.font.name = "Times New Roman"
    print(f"  ✅ Đã loại bỏ triệt để tên 'BabyBon' ({bb_count} đoạn văn được làm sạch).")

    # Sửa số lượng Use Case và Collection trong Kết luận & P147
    for p in doc.paragraphs:
        if "24 ca sử dụng (Use Case) chi tiết" in p.text:
            p.text = p.text.replace("24 ca sử dụng (Use Case) chi tiết", "27 ca sử dụng (Use Case) chi tiết (UC01 đến UC27)")
            for r in p.runs: r.font.name = "Times New Roman"
            print("  ✅ Đã sửa số lượng Use Case trong Kết luận: 27 ca sử dụng (UC01 đến UC27)")
        if "25 collection chuyên biệt" in p.text:
            p.text = p.text.replace("25 collection chuyên biệt", "31 collection (bao gồm Users)")
            for r in p.runs: r.font.name = "Times New Roman"
            print("  ✅ Đã sửa số lượng Collection trong Kết luận: 31 collection (bao gồm Users)")
        if "27 ca sử dụng cốt lõi (UC01 đến UC24)" in p.text:
            p.text = p.text.replace("27 ca sử dụng cốt lõi (UC01 đến UC24)", "27 ca sử dụng cốt lõi (UC01 đến UC27)")
            for r in p.runs: r.font.name = "Times New Roman"
            print("  ✅ Đã sửa typo tại P147: (UC01 đến UC27)")

    # =========================================================================
    # BƯỚC 2: CHUẨN HÓA MỤC 3.5 VÀ CẤU TRÚC ĐOẠN
    # =========================================================================
    print("\n--- BƯỚC 2: Chuẩn hóa Mục 3.5 và cấu trúc đoạn ---")
    p35 = None
    p_content = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("3.5. Sơ đồ"):
            p35 = p
            p_content = doc.paragraphs[i+1]
            break
            
    if p35 is not None:
        parent = p35._element.getparent()
        
        # Xóa các đoạn rỗng trước p35
        del_count = 0
        prev = p35._element.getprevious()
        while prev is not None and prev.tag.endswith('p'):
            prev_txt = docx.text.paragraph.Paragraph(prev, doc).text.strip()
            if not prev_txt and len(prev.xpath('.//w:drawing')) == 0:
                to_del = prev
                prev = prev.getprevious()
                parent.remove(to_del)
                del_count += 1
            else:
                break
        print(f"  ✅ Đã xóa {del_count} đoạn trống dư thừa trước Mục 3.5.")

        # Cập nhật style và tiêu đề Heading 2 cho Mục 3.5
        p35.style = doc.styles['Heading 2']
        p35.text = "3.5. Sơ đồ tổng quan mối liên hệ giữa các phân hệ nghiệp vụ kế toán và nguyên tắc hạch toán kép"
        p35.paragraph_format.keep_with_next = True
        p35.paragraph_format.space_before = Pt(12)
        p35.paragraph_format.space_after = Pt(6)
        for r in p35.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(14)
            r.font.bold = True

        # Trích xuất drawing từ p_content
        drawings = p_content._element.xpath('.//w:drawing')
        saved_drawing = drawings[0] if drawings else None
        print(f"  ✅ Trích xuất sơ đồ vector: {'Thành công' if saved_drawing is not None else 'Không tìm thấy'}")

        # Hàm helper chèn paragraph trước p_content
        def insert_p(text, style="Normal", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4, bullet=False):
            elem = OxmlElement('w:p')
            p = docx.text.paragraph.Paragraph(elem, doc)
            p.style = doc.styles[style] if style in doc.styles else doc.styles['Normal']
            p.alignment = align
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_after = Pt(space_after)
            r = p.add_run(text)
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            parent.insert(parent.index(p_content._element), elem)
            return p

        insert_p("Trong Hệ thống quản lý Cửa hàng Mẹ và Bé, các phân hệ nghiệp vụ được thiết kế liên kết hữu cơ theo dòng luân chuyển chứng từ kế toán chuẩn mực, đảm bảo mọi biến động hàng hóa vật chất đều được phản ánh đồng thời sang dòng tiền và số dư công nợ:")
        insert_p("1. Luồng Bán hàng – Xuất kho – Doanh thu & Công nợ: Khi phát sinh giao dịch bán lẻ tại quầy POS (DonHang), hệ thống tự động sinh Hóa đơn (HoaDon) và tạo Phiếu xuất kho (PhieuXuat), đồng thời cập nhật giảm số lượng tồn kho trong TonKho. Nếu khách hàng thanh toán ngay bằng tiền mặt hoặc chuyển khoản, hệ thống tự động tạo Phiếu thu (PhieuThu); nếu bán chịu, hệ thống tự động ghi nhận công nợ phải thu vào collection CongNo (LoaiCongNo: 'Khách hàng').")
        insert_p("2. Luồng Mua hàng – Nhập kho – Giá vốn & Công nợ NCC: Đơn đặt hàng (DonDatHang) gửi tới nhà cung cấp là căn cứ để tạo Phiếu nhập kho (PhieuNhap). Khi nhập kho, hệ thống tự động tăng tồn kho hàng hóa (TonKho), ghi nhận chi tiết dòng hàng vào CT_PhieuNhap và phát sinh công nợ phải trả vào collection CongNo (LoaiCongNo: 'Nhà cung cấp'). Khi thanh toán công nợ, kế toán lập Phiếu chi (PhieuChi) và đồng bộ giảm trừ công nợ.")
        insert_p("3. Nguyên tắc Hạch toán Nợ/Có tự động (On-The-Fly): Hệ thống áp dụng nguyên tắc tổng hợp động tại thời điểm truy vấn Sổ nhật ký chung (Mẫu S03a-DNN qua API GET /api/reports/general-journal) từ các chứng từ gốc theo đúng chế độ kế toán doanh nghiệp (Thông tư 133/2016/TT-BTC và Thông tư 200/2014/TT-BTC):")
        insert_p("• Nghiệp vụ Bán hàng: Nợ TK 111 (hoặc Nợ TK 131 nếu bán nợ) / Có TK 511 (Doanh thu bán hàng), Có TK 33311 (Thuế GTGT đầu ra); đồng thời ghi nhận Giá vốn: Nợ TK 632 / Có TK 156.", style="List Bullet", align=WD_ALIGN_PARAGRAPH.LEFT)
        insert_p("• Nghiệp vụ Mua hàng nhập kho: Nợ TK 156 (Hàng hóa) / Có TK 331 (Phải trả NCC) hoặc Có TK 111 (nếu thanh toán ngay).", style="List Bullet", align=WD_ALIGN_PARAGRAPH.LEFT)
        insert_p("• Nghiệp vụ Thu tiền mặt: Nợ TK 111 / Có TK 131 (Thu nợ khách hàng).", style="List Bullet", align=WD_ALIGN_PARAGRAPH.LEFT)
        insert_p("• Nghiệp vụ Chi tiền mặt: Nợ TK 331 (Thanh toán nợ NCC) / Có TK 111.", style="List Bullet", align=WD_ALIGN_PARAGRAPH.LEFT)
        insert_p("*(Lưu ý: Các giao dịch thanh toán chuyển khoản qua ngân hàng tại quầy được hạch toán đồng bộ về dòng tiền quỹ tương đương tiền. Nghiệp vụ kiểm kê kho và đổi trả hàng được xử lý cân bằng kho và hoàn tiền trực tiếp trên chứng từ)*.")

        # Đoạn chứa ảnh sơ đồ
        img_elem = OxmlElement('w:p')
        img_p = docx.text.paragraph.Paragraph(img_elem, doc)
        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p.paragraph_format.space_before = Pt(6)
        img_p.paragraph_format.space_after = Pt(4)
        if saved_drawing is not None:
            run_elem = OxmlElement('w:r')
            run_elem.append(saved_drawing)
            img_elem.append(run_elem)
        else:
            r = img_p.add_run("[CHÈN HÌNH: Sơ đồ luân chuyển chứng từ và mối liên hệ giữa các phân hệ nghiệp vụ kế toán]")
            r.font.name = "Times New Roman"
            r.font.italic = True
        parent.insert(parent.index(p_content._element), img_elem)

        # Đoạn chú thích Hình 3.2
        cap_elem = OxmlElement('w:p')
        cap_p = docx.text.paragraph.Paragraph(cap_elem, doc)
        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_p.paragraph_format.space_after = Pt(8)
        cap_p.paragraph_format.keep_with_next = True
        r1 = cap_p.add_run("Hình 3.2. ")
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(10.5)
        r1.font.bold = True
        r1.font.italic = True
        r2 = cap_p.add_run("Sơ đồ luân chuyển chứng từ và mối liên hệ giữa các phân hệ nghiệp vụ kế toán")
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(10.5)
        r2.font.italic = True
        parent.insert(parent.index(p_content._element), cap_elem)

        # Đoạn kết luận chuyển tiếp sang Chương 4
        insert_p("Mô hình liên kết chứng từ và cơ chế hạch toán tự động nêu trên là nền tảng cốt lõi kết nối giữa tầng thiết kế cơ sở dữ liệu (Chương 3) và tầng giao diện tương tác người dùng (Chương 4). Thông qua cấu trúc 31 collection được chuẩn hóa và các ràng buộc toàn vẹn nghiêm ngặt, hệ thống bảo đảm tính chính xác của số liệu kế toán tại mọi điểm chạm nghiệp vụ, sẵn sàng đáp ứng yêu cầu vận hành của các màn hình chức năng được trình bày chi tiết trong chương tiếp theo.", space_after=8)

        # Xóa đoạn cũ p_content
        parent.remove(p_content._element)
        print("  ✅ Đã tái cấu trúc hoàn chỉnh Mục 3.5 và kết nối logic sang Chương 4.")

    # =========================================================================
    # BƯỚC 3: CHUẨN HÓA ĐÁNH SỐ HÌNH ẢNH VÀ PLACEHOLDERS (2.1 & 2.5)
    # =========================================================================
    print("\n--- BƯỚC 3: Chuẩn hóa đánh số hình ảnh và Placeholders ---")
    figure_caption_map = {
        "Hình 2.1.1. Sơ đồ quy trình mua hàng": "Hình 2.1. Sơ đồ quy trình mua hàng",
        "Hình 2.2.1. Sơ đồ quy trình bán hàng": "Hình 2.2. Sơ đồ quy trình bán hàng",
        "Hình 2.3.1. Sơ đồ quy trình nhập kho": "Hình 2.3. Sơ đồ quy trình nhập kho",
        "Hình 2.4.1. Sơ đồ quy trình xuất kho": "Hình 2.4. Sơ đồ quy trình xuất kho",
        "Hình 2.5.1. Sơ đồ quy trình thu tiền từ khách hàng": "Hình 2.5. Sơ đồ quy trình thu tiền từ khách hàng",
        "Hình 2.5.2. Sơ đồ quy trình chi tiền cho nhà cung cấp": "Hình 2.6. Sơ đồ quy trình chi tiền cho nhà cung cấp",
        "Hình 2.4.1. Biểu đồ Use Case tổng thể hệ thống": "Hình 2.7. Biểu đồ Use Case tổng thể hệ thống",
        "Hình 2.4.2. Biểu đồ Use Case phân hệ Quản trị và Quản lý danh mục": "Hình 2.8. Biểu đồ Use Case phân hệ Quản trị và Quản lý danh mục",
        "Hình 2.4.3. Biểu đồ Use Case phân hệ Mua hàng và Nhập kho": "Hình 2.9. Biểu đồ Use Case phân hệ Mua hàng và Nhập kho",
        "Hình 2.4.4. Biểu đồ Use Case phân hệ Bán hàng và Xuất kho": "Hình 2.10. Biểu đồ Use Case phân hệ Bán hàng và Xuất kho",
        "Hình 2.4.5. Biểu đồ Use Case phân hệ Kế toán Thu chi và Công nợ": "Hình 2.11. Biểu đồ Use Case phân hệ Kế toán Thu chi và Công nợ",
        "Hình 2.4.6. Biểu đồ Use Case phân hệ Quản lý Kho và Báo cáo điều hành": "Hình 2.12. Biểu đồ Use Case phân hệ Quản lý Kho và Báo cáo điều hành",
        "Hình 2.5.7. Biểu đồ hoạt động quy trình Mua hàng": "Hình 2.13. Biểu đồ hoạt động quy trình Mua hàng",
        "Hình 2.5.8. Biểu đồ hoạt động quy trình Bán hàng": "Hình 2.14. Biểu đồ hoạt động quy trình Bán hàng",
        "Hình 2.5.9. Biểu đồ hoạt động quy trình Nhập kho": "Hình 2.15. Biểu đồ hoạt động quy trình Nhập kho",
        "Hình 2.5.10. Biểu đồ hoạt động quy trình Xuất kho": "Hình 2.16. Biểu đồ hoạt động quy trình Xuất kho",
        "Hình 2.5.11. Biểu đồ hoạt động quy trình Thu tiền từ Khách hàng": "Hình 2.17. Biểu đồ hoạt động quy trình Thu tiền từ Khách hàng",
        "Hình 2.5.12. Biểu đồ hoạt động quy trình Chi tiền cho Nhà cung cấp": "Hình 2.18. Biểu đồ hoạt động quy trình Chi tiền cho Nhà cung cấp",
        "Hình 3.3.1. Sơ đồ quan hệ logic giữa các collection": "Hình 3.1. Sơ đồ quan hệ logic giữa các collection (ERD)",
        "Hình 4.2.1. Giao diện màn hình Đăng nhập": "Hình 4.1. Giao diện màn hình Đăng nhập",
        "Hình 4.2.2. Giao diện màn hình Tổng quan": "Hình 4.2. Giao diện màn hình Tổng quan (Dashboard)",
        "Hình 4.2.3. Giao diện màn hình Quản lý Sản phẩm": "Hình 4.3. Giao diện màn hình Quản lý Sản phẩm",
        "Hình 4.2.4. Giao diện màn hình Quản lý Loại hàng": "Hình 4.4. Giao diện màn hình Quản lý Loại hàng",
        "Hình 4.2.5. Giao diện màn hình Quản lý Khách hàng": "Hình 4.5. Giao diện màn hình Quản lý Khách hàng",
        "Hình 4.2.6. Giao diện màn hình Quản lý Nhà cung cấp": "Hình 4.6. Giao diện màn hình Quản lý Nhà cung cấp",
        "Hình 4.2.7. Giao diện màn hình Quản lý Nhân viên": "Hình 4.7. Giao diện màn hình Quản lý Nhân viên",
        "Hình 4.2.8. Giao diện màn hình Lập đơn mua hàng": "Hình 4.8. Giao diện màn hình Lập đơn mua hàng",
        "Hình 4.2.9. Giao diện màn hình Phiếu nhập kho": "Hình 4.9. Giao diện màn hình Phiếu nhập kho",
        "Hình 4.2.10. Giao diện màn hình Hóa đơn bán hàng": "Hình 4.10. Giao diện màn hình Hóa đơn bán hàng",
        "Hình 4.2.11. Giao diện màn hình Phiếu xuất kho": "Hình 4.11. Giao diện màn hình Phiếu xuất kho",
        "Hình 4.2.12. Giao diện màn hình Phiếu thu": "Hình 4.12. Giao diện màn hình Phiếu thu",
        "Hình 4.2.13. Giao diện màn hình Phiếu chi": "Hình 4.13. Giao diện màn hình Phiếu chi",
        "Hình 4.2.14. Giao diện màn hình Quản lý tồn kho": "Hình 4.14. Giao diện màn hình Quản lý tồn kho",
        "Hình 4.2.15. Giao diện màn hình Kiểm kê kho": "Hình 4.15. Giao diện màn hình Kiểm kê kho",
        "Hình 4.2.16. Giao diện màn hình Báo cáo": "Hình 4.16. Giao diện màn hình Hệ thống 17 Báo cáo Kế toán & Quản trị",
        "Hình 4.2.17. Giao diện màn hình Đổi trả hàng": "Hình 4.17. Giao diện màn hình Đổi trả hàng (ReturnPage)",
        "Hình 4.2.18. Giao diện màn hình Quản lý Công nợ Nhà cung cấp": "Hình 4.18. Giao diện màn hình Quản lý Công nợ Nhà cung cấp (DebtsPage)",
        "Hình 4.2.19. Giao diện màn hình Quản lý Khuyến mại": "Hình 4.19. Giao diện màn hình Quản lý Khuyến mại & Mã giảm giá (PromotionsPage)",
        "Hình 4.2.20. Giao diện màn hình Nhật ký kiểm toán và Sao lưu dữ liệu": "Hình 4.20. Giao diện màn hình Quản trị hệ thống: Nhật ký kiểm toán & Sao lưu dữ liệu"
    }

    fig_updated = 0
    for p in doc.paragraphs:
        txt = p.text.strip()
        for old_cap, new_cap in figure_caption_map.items():
            if txt == old_cap or txt.startswith(old_cap):
                prefix = new_cap.split(". ")[0] + ". "
                title = new_cap[len(prefix):]
                p.text = ""
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(6)
                
                r1 = p.add_run(prefix)
                r1.font.name = "Times New Roman"
                r1.font.size = Pt(10.5)
                r1.font.bold = True
                r1.font.italic = True
                
                r2 = p.add_run(title)
                r2.font.name = "Times New Roman"
                r2.font.size = Pt(10.5)
                r2.font.bold = False
                r2.font.italic = True
                fig_updated += 1
                break
    print(f"  ✅ Đã chuẩn hóa {fig_updated} chú thích hình ảnh (Hình 2.1-2.18, Hình 3.1-3.2, Hình 4.1-4.20).")

    # Chuẩn hóa Placeholder format: [CHÈN HÌNH: ...]
    ph_count = 0
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith("⟦") and txt.endswith("⟧"):
            clean_content = txt[1:-1].strip()
            if clean_content.startswith("Ảnh chụp minh họa thực tế:"):
                clean_content = clean_content.replace("Ảnh chụp minh họa thực tế:", "").strip()
            new_ph = f"[CHÈN HÌNH: {clean_content}]"
            p.text = new_ph
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11)
                r.font.italic = True
                r.font.color.rgb = RGBColor(100, 116, 139)
            ph_count += 1
    print(f"  ✅ Đã chuẩn hóa {ph_count} placeholder hình ảnh dạng [CHÈN HÌNH: ...].")

    # =========================================================================
    # BƯỚC 4: CHUẨN HÓA ĐÁNH SỐ BẢNG VÀ CHÚ THÍCH BẢNG (2.1)
    # =========================================================================
    print("\n--- BƯỚC 4: Chuẩn hóa đánh số bảng và chú thích bảng ---")
    table_caption_map = {
        "Bảng 2.1.1. Bảng mô tả các bước trong quy trình mua hàng": "Bảng 2.1. Mô tả các bước trong quy trình mua hàng",
        "Bảng 2.1.2. Bảng mô tả các bước trong quy trình bán hàng": "Bảng 2.2. Mô tả các bước trong quy trình bán hàng",
        "Bảng 2.1.3. Bảng mô tả các bước trong quy trình nhập kho": "Bảng 2.3. Mô tả các bước trong quy trình nhập kho",
        "Bảng 2.1.4. Bảng mô tả các bước trong quy trình xuất kho": "Bảng 2.4. Mô tả các bước trong quy trình xuất kho",
        "Bảng 2.1.5. Bảng mô tả các bước thu tiền từ khách hàng": "Bảng 2.5. Mô tả các bước thu tiền từ khách hàng",
        "Bảng 2.1.6. Bảng mô tả các bước chi tiền cho nhà cung cấp": "Bảng 2.6. Mô tả các bước chi tiền cho nhà cung cấp",
        "Bảng 2.2.1. Danh sách các tác nhân và ma trận phân quyền trong hệ thống": "Bảng 2.7. Danh sách các tác nhân và ma trận phân quyền trong hệ thống",
        "Bảng 2.3.1. Danh mục 27 Use Case của Hệ thống quản lý Cửa hàng Mẹ và Bé": "Bảng 2.8. Danh mục 27 Use Case của Hệ thống quản lý Cửa hàng Mẹ và Bé"
    }

    tbl_updated = 0
    for p in doc.paragraphs:
        txt = p.text.strip()
        for old_cap, new_cap in table_caption_map.items():
            if txt == old_cap or txt.startswith(old_cap):
                prefix = new_cap.split(". ")[0] + ". "
                title = new_cap[len(prefix):]
                p.text = ""
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(3)
                
                r1 = p.add_run(prefix)
                r1.font.name = "Times New Roman"
                r1.font.size = Pt(10.5)
                r1.font.bold = True
                
                r2 = p.add_run(title)
                r2.font.name = "Times New Roman"
                r2.font.size = Pt(10.5)
                r2.font.bold = True
                tbl_updated += 1
                break
    print(f"  ✅ Đã chuẩn hóa {tbl_updated} chú thích bảng Chương 2 (Bảng 2.1 đến Bảng 2.8).")

    collection_table_names = [
        "Users", "VaiTro", "NhanVien", "KhachHang", "NhaCungCap",
        "LoaiHang", "SanPham", "DonDatHang", "CT_DonDatHang", "DonHang",
        "CT_DonHang", "HoaDon", "CT_HoaDon", "PhieuNhap", "CT_PhieuNhap",
        "PhieuXuat", "CT_PhieuXuat", "TonKho", "KiemKe", "CT_KiemKe",
        "DieuChinhKho", "PhieuTraHang", "CT_PhieuTraHang", "PhieuThu", "PhieuChi",
        "ThanhToan", "CongNo", "KhuyenMai", "CT_KhuyenMai", "Counters", "AuditLogs"
    ]

    body_elements = doc._body._body
    t_idx = 0
    c_tbl_added = 0
    
    for child in list(body_elements):
        if child.tag.endswith('tbl'):
            if 9 <= t_idx <= 39:
                c_num = t_idx - 8
                col_name = collection_table_names[c_num - 1]
                
                prev = child.getprevious()
                while prev is not None and not prev.tag.endswith('p'):
                    prev = prev.getprevious()
                
                prev_text = ""
                if prev is not None:
                    prev_p = docx.text.paragraph.Paragraph(prev, doc)
                    prev_text = prev_p.text.strip()
                    
                if not prev_text.startswith(f"Bảng 3.{c_num}."):
                    cap_elem = OxmlElement('w:p')
                    cap_p = docx.text.paragraph.Paragraph(cap_elem, doc)
                    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap_p.paragraph_format.space_before = Pt(8)
                    cap_p.paragraph_format.space_after = Pt(3)
                    cap_p.paragraph_format.keep_with_next = True
                    
                    prefix = f"Bảng 3.{c_num}. "
                    r1 = cap_p.add_run(prefix)
                    r1.font.name = "Times New Roman"
                    r1.font.size = Pt(10.5)
                    r1.font.bold = True
                    
                    r2 = cap_p.add_run(f"Cấu trúc dữ liệu collection {col_name}")
                    r2.font.name = "Times New Roman"
                    r2.font.size = Pt(10.5)
                    r2.font.bold = True
                    
                    body_elements.insert(body_elements.index(child), cap_elem)
                    c_tbl_added += 1
            t_idx += 1
            
    print(f"  ✅ Đã gắn chú thích Bảng 3.1 đến Bảng 3.31 cho toàn bộ 31 Collection ({c_tbl_added} bảng được bổ sung chú thích).")

    # =========================================================================
    # BƯỚC 5: ĐỊNH DẠNG BẢNG BIỂU CHUẨN MỰC (TBLHEADER, CANTSPLIT, CĂN GIỮA)
    # =========================================================================
    print("\n--- BƯỚC 5: Chuẩn hóa định dạng kỹ thuật bảng biểu ---")
    for tbl in doc.tables:
        tblPr = tbl._tbl.tblPr
        for jc in tblPr.xpath('.//w:jc'):
            tblPr.remove(jc)
        tblPr.append(parse_xml(f'<w:jc {nsdecls("w")} w:val="center"/>'))
        
        if len(tbl.rows) > 0:
            header_trPr = tbl.rows[0]._tr.get_or_add_trPr()
            if len(header_trPr.xpath('.//w:tblHeader')) == 0:
                header_trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            if len(header_trPr.xpath('.//w:cantSplit')) == 0:
                header_trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
                
            for cell in tbl.rows[0].cells:
                set_cell_background(cell, "F1F5F9")
                set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(9.5)
                        r.font.bold = True
                        
        for row in tbl.rows[1:]:
            trPr = row._tr.get_or_add_trPr()
            if len(trPr.xpath('.//w:cantSplit')) == 0:
                trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            for cell in row.cells:
                set_cell_margins(cell, top=70, bottom=70, left=120, right=120)
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(9.5)

    print(f"  ✅ Đã áp dụng tblHeader, cantSplit và căn giữa cho toàn bộ {len(doc.tables)} bảng.")

    # =========================================================================
    # BƯỚC 6: CHUẨN HÓA TYPOGRAPHY, HEADINGS VÀ DỌN DẸP ĐOẠN TRỐNG (2.6)
    # =========================================================================
    print("\n--- BƯỚC 6: Chuẩn hóa Tiêu đề & Loại bỏ đoạn trống dư thừa ---")
    headings_fixed = 0
    for p in doc.paragraphs:
        if p.style.name.startswith("Heading"):
            p.paragraph_format.keep_with_next = True
            headings_fixed += 1
            for r in p.runs:
                r.font.name = "Times New Roman"
    print(f"  ✅ Đã thiết lập keep_with_next = True cho {headings_fixed} tiêu đề các cấp.")

    # Dọn dẹp các đoạn trống dư thừa liên tiếp
    cleaned_empty = 0
    consecutive_empty = 0
    for p in list(doc.paragraphs):
        is_empty = (not p.text.strip()) and (len(p._element.xpath('.//w:drawing')) == 0) and (len(p._element.xpath('.//w:br[@w:type="page"]')) == 0)
        if is_empty:
            consecutive_empty += 1
            if consecutive_empty > 1:
                p._element.getparent().remove(p._element)
                cleaned_empty += 1
        else:
            consecutive_empty = 0
    print(f"  ✅ Đã dọn dẹp {cleaned_empty} đoạn trống dư thừa liên tiếp.")

    # =========================================================================
    # BƯỚC 7: THIẾT LẬP FRONT MATTER (MỤC LỤC, DANH MỤC HÌNH, DANH MỤC BẢNG)
    # =========================================================================
    print("\n--- BƯỚC 7: Thiết lập Mục lục & Danh mục Hình, Bảng tại Front Matter ---")
    
    insert_point_p = None
    for p in doc.paragraphs:
        if p.text.strip() == "DANH MỤC CÁC TỪ VIẾT TẮT":
            insert_point_p = p
            break
            
    if insert_point_p is not None:
        parent = insert_point_p._element.getparent()
        
        def insert_p_before(text="", style="Normal", align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, italic=False, size=12, before=0, after=4, keep_next=False, page_break_after=False):
            p_elem = OxmlElement('w:p')
            p_obj = docx.text.paragraph.Paragraph(p_elem, doc)
            p_obj.style = doc.styles[style] if style in doc.styles else doc.styles['Normal']
            p_obj.alignment = align
            p_obj.paragraph_format.space_before = Pt(before)
            p_obj.paragraph_format.space_after = Pt(after)
            p_obj.paragraph_format.keep_with_next = keep_next
            p_obj.paragraph_format.line_spacing = 1.3
            if text:
                r = p_obj.add_run(text)
                r.font.name = "Times New Roman"
                r.font.size = Pt(size)
                r.font.bold = bold
                r.font.italic = italic
            if page_break_after:
                br_elem = OxmlElement('w:r')
                br = parse_xml(f'<w:br {nsdecls("w")} w:type="page"/>')
                br_elem.append(br)
                p_elem.append(br_elem)
            parent.insert(parent.index(insert_point_p._element), p_elem)
            return p_obj
            
        # 1. Trang MỤC LỤC
        insert_p_before("MỤC LỤC", style="Heading 1", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16, before=12, after=12, keep_next=True)
        
        toc_p = OxmlElement('w:p')
        toc_fld = parse_xml(f'''
            <w:fldSimple {nsdecls("w")} w:instr="TOC \\o &quot;1-3&quot; \\h \\z \\u">
                <w:r>
                    <w:rPr><w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/><w:sz w:val="24"/></w:rPr>
                    <w:t>[Mục lục tự động sẽ được Microsoft Word cập nhật khi mở tài liệu]</w:t>
                </w:r>
            </w:fldSimple>
        ''')
        toc_p.append(toc_fld)
        parent.insert(parent.index(insert_point_p._element), toc_p)
        insert_p_before(page_break_after=True)
        
        # 2. Trang DANH MỤC HÌNH ẢNH
        insert_p_before("DANH MỤC HÌNH ẢNH", style="Heading 1", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16, before=12, after=12, keep_next=True)
        
        # Danh sách 40 hình theo đúng thứ tự
        ordered_figs = [
            "Hình 2.1. Sơ đồ quy trình mua hàng",
            "Hình 2.2. Sơ đồ quy trình bán hàng",
            "Hình 2.3. Sơ đồ quy trình nhập kho",
            "Hình 2.4. Sơ đồ quy trình xuất kho",
            "Hình 2.5. Sơ đồ quy trình thu tiền từ khách hàng",
            "Hình 2.6. Sơ đồ quy trình chi tiền cho nhà cung cấp",
            "Hình 2.7. Biểu đồ Use Case tổng thể hệ thống",
            "Hình 2.8. Biểu đồ Use Case phân hệ Quản trị và Quản lý danh mục",
            "Hình 2.9. Biểu đồ Use Case phân hệ Mua hàng và Nhập kho",
            "Hình 2.10. Biểu đồ Use Case phân hệ Bán hàng và Xuất kho",
            "Hình 2.11. Biểu đồ Use Case phân hệ Kế toán Thu chi và Công nợ",
            "Hình 2.12. Biểu đồ Use Case phân hệ Quản lý Kho và Báo cáo điều hành",
            "Hình 2.13. Biểu đồ hoạt động quy trình Mua hàng",
            "Hình 2.14. Biểu đồ hoạt động quy trình Bán hàng",
            "Hình 2.15. Biểu đồ hoạt động quy trình Nhập kho",
            "Hình 2.16. Biểu đồ hoạt động quy trình Xuất kho",
            "Hình 2.17. Biểu đồ hoạt động quy trình Thu tiền từ Khách hàng",
            "Hình 2.18. Biểu đồ hoạt động quy trình Chi tiền cho Nhà cung cấp",
            "Hình 3.1. Sơ đồ quan hệ logic giữa các collection (ERD)",
            "Hình 3.2. Sơ đồ luân chuyển chứng từ và mối liên hệ giữa các phân hệ nghiệp vụ kế toán",
            "Hình 4.1. Giao diện màn hình Đăng nhập",
            "Hình 4.2. Giao diện màn hình Tổng quan (Dashboard)",
            "Hình 4.3. Giao diện màn hình Quản lý Sản phẩm",
            "Hình 4.4. Giao diện màn hình Quản lý Loại hàng",
            "Hình 4.5. Giao diện màn hình Quản lý Khách hàng",
            "Hình 4.6. Giao diện màn hình Quản lý Nhà cung cấp",
            "Hình 4.7. Giao diện màn hình Quản lý Nhân viên",
            "Hình 4.8. Giao diện màn hình Lập đơn mua hàng",
            "Hình 4.9. Giao diện màn hình Phiếu nhập kho",
            "Hình 4.10. Giao diện màn hình Hóa đơn bán hàng",
            "Hình 4.11. Giao diện màn hình Phiếu xuất kho",
            "Hình 4.12. Giao diện màn hình Phiếu thu",
            "Hình 4.13. Giao diện màn hình Phiếu chi",
            "Hình 4.14. Giao diện màn hình Quản lý tồn kho",
            "Hình 4.15. Giao diện màn hình Kiểm kê kho",
            "Hình 4.16. Giao diện màn hình Hệ thống 17 Báo cáo Kế toán & Quản trị",
            "Hình 4.17. Giao diện màn hình Đổi trả hàng (ReturnPage)",
            "Hình 4.18. Giao diện màn hình Quản lý Công nợ Nhà cung cấp (DebtsPage)",
            "Hình 4.19. Giao diện màn hình Quản lý Khuyến mại & Mã giảm giá (PromotionsPage)",
            "Hình 4.20. Giao diện màn hình Quản trị hệ thống: Nhật ký kiểm toán & Sao lưu dữ liệu"
        ]
        for fig_line in ordered_figs:
            insert_p_before(fig_line, style="Normal", size=11, before=1, after=2)
        insert_p_before(page_break_after=True)

        # 3. Trang DANH MỤC BẢNG BIỂU
        insert_p_before("DANH MỤC BẢNG BIỂU", style="Heading 1", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16, before=12, after=12, keep_next=True)
        ordered_tbls = [
            "Bảng 2.1. Mô tả các bước trong quy trình mua hàng",
            "Bảng 2.2. Mô tả các bước trong quy trình bán hàng",
            "Bảng 2.3. Mô tả các bước trong quy trình nhập kho",
            "Bảng 2.4. Mô tả các bước trong quy trình xuất kho",
            "Bảng 2.5. Mô tả các bước thu tiền từ khách hàng",
            "Bảng 2.6. Mô tả các bước chi tiền cho nhà cung cấp",
            "Bảng 2.7. Danh sách các tác nhân và ma trận phân quyền trong hệ thống",
            "Bảng 2.8. Danh mục 27 Use Case của Hệ thống quản lý Cửa hàng Mẹ và Bé"
        ]
        for c_idx, c_name in enumerate(collection_table_names, 1):
            ordered_tbls.append(f"Bảng 3.{c_idx}. Cấu trúc dữ liệu collection {c_name}")
        ordered_tbls.extend([
            "Bảng A.1. Phân chia module hệ thống theo đặc tả chức năng",
            "Bảng A.2. Phân công nhiệm vụ các thành viên trong nhóm",
            "Bảng A.3. Kế hoạch và tiến độ thực hiện đề tài",
            "Bảng B.1. Nhật ký đánh giá kết quả thực hiện của từng thành viên"
        ])
        for tbl_line in ordered_tbls:
            insert_p_before(tbl_line, style="Normal", size=11, before=1, after=2)
        insert_p_before(page_break_after=True)
        
        print("  ✅ Đã chèn Mục lục tự động, Danh mục hình ảnh và Danh mục bảng biểu vào Front Matter.")

    # =========================================================================
    # BƯỚC 8: CẤU HÌNH UPDATE FIELDS TRONG SETTINGS.XML & LƯU FILE
    # =========================================================================
    print(f"\n--- BƯỚC 8: Lưu file {dst_docx} ---")
    settingsPr = doc.settings.element
    if len(settingsPr.xpath('.//w:updateFields')) == 0:
        settingsPr.append(parse_xml(f'<w:updateFields {nsdecls("w")} w:val="true"/>'))
        
    doc.save(dst_docx)
    print(f"  ✅ Đã lưu thành công: {dst_docx}")

    # =========================================================================
    # BƯỚC 9: XUẤT FILE PDF QUA MICROSOFT WORD COM
    # =========================================================================
    print(f"\n--- BƯỚC 9: Cập nhật trường & Xuất PDF ({dst_pdf}) ---")
    try:
        import win32com.client
        cwd = os.getcwd()
        full_docx = os.path.join(cwd, dst_docx)
        full_pdf = os.path.join(cwd, dst_pdf)
        
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        try:
            print("  Đang mở file qua Word Application...")
            w_doc = word.Documents.Open(full_docx)
            print("  Đang cập nhật toàn bộ trường (Fields & TOC)...")
            w_doc.Fields.Update()
            w_doc.Save() # Lưu lại sau khi cập nhật TOC
            print(f"  Đang kết xuất sang PDF: {dst_pdf}...")
            w_doc.SaveAs(full_pdf, FileFormat=17) # 17 = wdFormatPDF
            w_doc.Close()
            pdf_size = os.path.getsize(full_pdf)
            print(f"  ✅ ĐÃ XUẤT THÀNH CÔNG: {dst_pdf} (Dung lượng: {pdf_size:,} bytes)")
        finally:
            word.Quit()
    except Exception as e:
        print(f"  ⚠️ Cảnh báo xuất PDF: {e}")
        print("  Vẫn bảo đảm file DOCX hoàn chỉnh tại baocaobtl_final.docx")

    print("\n===============================================================")
    print("HOÀN THÀNH TOÀN BỘ CÁC HẠNG MỤC PHA 3!")
    print("===============================================================")

if __name__ == "__main__":
    main()
