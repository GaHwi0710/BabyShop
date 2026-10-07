import docx
import sys
import re

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('baocaobtl_final.docx')
current_h1 = ''
current_h2 = ''
current_h3 = ''
placeholders = []

for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if p.style.name == 'Heading 1':
        current_h1 = txt
    elif p.style.name == 'Heading 2':
        current_h2 = txt
    elif p.style.name == 'Heading 3':
        current_h3 = txt
    if txt.startswith('[CHÈN HÌNH:'):
        caption = ''
        if i + 1 < len(doc.paragraphs) and re.match(r'^Hình\s+[0-9]+\.[0-9]+', doc.paragraphs[i+1].text.strip()):
            caption = doc.paragraphs[i+1].text.strip()
        placeholders.append({
            'index': len(placeholders) + 1,
            'p_idx': i,
            'h1': current_h1,
            'h2': current_h2,
            'h3': current_h3,
            'placeholder': txt,
            'caption': caption
        })

for item in placeholders:
    print(f"{item['index']}|{item['h1']}|{item['h2']}|{item['h3']}|{item['caption']}|{item['placeholder']}")
