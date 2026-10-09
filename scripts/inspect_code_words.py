# -*- coding: utf-8 -*-
import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document('D:/BabyShop/Nhom1_BCBTL.docx')

print("--- Courier New in Paragraphs ---")
courier_p_count = 0
for i, p in enumerate(doc.paragraphs):
    for r in p.runs:
        if r.font.name == 'Courier New':
            courier_p_count += 1
            if courier_p_count <= 25:
                print(f"P{i} (size={r.font.size.pt if r.font.size else 'None'}): '{r.text}'")
print(f"Total Courier New runs in paragraphs: {courier_p_count}")

print("\n--- Courier New in Tables ---")
courier_t_count = 0
for ti, t in enumerate(doc.tables):
    for ri, row in enumerate(t.rows):
        for ci, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                for r in p.runs:
                    if r.font.name == 'Courier New':
                        courier_t_count += 1
                        if courier_t_count <= 25:
                            print(f"T{ti} R{ri} C{ci} (size={r.font.size.pt if r.font.size else 'None'}): '{r.text}'")
print(f"Total Courier New runs in tables: {courier_t_count}")

# Check what other fonts exist in oxml w:rFonts
print("\n--- XML rFonts check ---")
rfonts_set = set()
for p in doc.paragraphs:
    for r in p.runs:
        rFonts = r._r.xpath('.//w:rFonts')
        for rf in rFonts:
            for k in rf.attrib:
                rfonts_set.add((k, rf.attrib[k]))
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    rFonts = r._r.xpath('.//w:rFonts')
                    for rf in rFonts:
                        for k in rf.attrib:
                            rfonts_set.add((k, rf.attrib[k]))
print("All XML rFonts attributes:", rfonts_set)
