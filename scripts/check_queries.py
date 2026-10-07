import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('baocaobtl_fixed.docx')
queries = [
    '25 collection', '24 Use Case', '24 ca sử dụng', '24 UC', '16 màn hình', 
    '5 nhóm tác nhân', 'Schema Validation khắt khe', 'mã vạch', 'barcode',
    'inventory_transactions', 'pending'
]

for q in queries:
    p_matches = [(i, doc.paragraphs[i].text[:80]) for i, p in enumerate(doc.paragraphs) if q.lower() in p.text.lower()]
    print(f"Query '{q}': found {len(p_matches)}")
    for idx, txt in p_matches:
        print(f"   P{idx}: {txt}")
