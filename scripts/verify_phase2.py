import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('baocaobtl_fixed.docx')
print("=== VERIFY PHASE 2 AUDIT ===")
print("Total paragraphs:", len(doc.paragraphs))
print("Total tables:", len(doc.tables))

# 1. Check "31 collection (bao gồm Users)"
p_31_count = sum(1 for p in doc.paragraphs if "31 collection (bao gồm Users)" in p.text)
print(f"1. '31 collection (bao gồm Users)': {p_31_count} (expect >= 2)")

# 2. Check old collection names in paragraphs 150-550
old_names = [
    '`users`', '`categories`', '`employees`', '`products`', '`suppliers`', 
    '`inventory`', '`purchase_orders`', '`warehouse_receipts`', 
    '`warehouse_exports`', '`sales_invoices`', '`cash_receipts`', 
    '`cash_payments`', '`stock_takes`', 'inventory_transactions', "'pending'"
]
uc_text = " ".join(doc.paragraphs[i].text for i in range(150, min(550, len(doc.paragraphs))))
print("2. Check old collection names / artifacts in UCs:")
for name in old_names:
    c = uc_text.count(name)
    print(f"   {name}: {c}")

# 3. Check old vs new business codes in UCs
old_codes = ['KH-00001', 'NCC-00001', 'DH-00001', 'PN-00001', 'PX-00001', 'PT-00001', 'PC-00001', 'HD-00001', 'KK-00001']
print("3. Check old business codes in UCs:")
for code in old_codes:
    c = uc_text.count(code)
    print(f"   {code}: {c} (expect 0)")

new_codes = ['KH001', 'NCC001', 'DDH001', 'PN001', 'PX001', 'PT001', 'PC001', 'HD001', 'KK001']
print("4. Check new business codes in UCs:")
for code in new_codes:
    c = uc_text.count(code)
    print(f"   {code}: {c} (expect >= 1)")

# 4. Check UC25, UC26, UC27
print("5. Check UC25, UC26, UC27 in text:")
all_text = " ".join(p.text for p in doc.paragraphs)
for uc in ['UC25 – Quản lý đổi trả hàng', 'UC26 – Quản lý tích điểm và đổi voucher', 'UC27 – Quản trị sao lưu và phục hồi dữ liệu']:
    print(f"   '{uc}': {all_text.count(uc)}")

# 5. Check Table 8 (Use Cases)
t8 = doc.tables[8]
print(f"6. Table 8 total rows: {len(t8.rows)} (expect 28: 1 header + 27 UCs)")
codes_in_t8 = [r.cells[0].text.strip() for r in t8.rows[1:]]
print(f"   First 3: {codes_in_t8[:3]}, Last 3: {codes_in_t8[-3:]}")

# 6. Check Barcode / Mã vạch occurrences
print("7. Check 'mã vạch' and 'barcode' in document:")
for i, p in enumerate(doc.paragraphs):
    if 'mã vạch' in p.text.lower() or 'barcode' in p.text.lower():
        print(f"   P{i}: {p.text[:90]}")
for t_idx, t in enumerate(doc.tables):
    for r_idx, r in enumerate(t.rows):
        for c_idx, c in enumerate(r.cells):
            if 'mã vạch' in c.text.lower() or 'barcode' in c.text.lower():
                print(f"   Table {t_idx} R{r_idx} C{c_idx}: {c.text[:80]}")

# 7. Check Section 3.5
p_35 = [p.text[:80] for p in doc.paragraphs if "3.5. Sơ đồ tổng quan" in p.text]
print("8. Section 3.5 heading:", p_35)

# 8. Check Chapter 4 screens (4.2.1 to 4.2.20)
screens_42 = [p.text.strip() for p in doc.paragraphs if p.text.strip().startswith("4.2.") and ("Màn hình" in p.text or "Hệ thống" in p.text)]
print(f"9. Total screens in 4.2: {len(screens_42)}")
for s in screens_42:
    print(f"   {s}")

# 9. Check Report #15 disclaimer
r15_check = sum(1 for p in doc.paragraphs if "isModuleAvailable: false" in p.text or "chưa có thuật toán tính khấu hao" in p.text)
print(f"10. Report #15 disclaimer in text: {r15_check} (expect >= 1)")

# 10. Check Conclusion & Appendix B.1
print("11. Check Conclusion & Appendix B.1:")
for p in doc.paragraphs[850:900]:
    if any(k in p.text for k in ['27 ca sử dụng', '31 collection', 'VietQR', 'khấu hao', 'mã vạch']):
        print(f"   P: {p.text[:90]}")
t43 = doc.tables[43]
print(f"   Table 43 R6 C3: {t43.rows[6].cells[3].text}")
