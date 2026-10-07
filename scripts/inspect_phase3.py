import docx
import sys
import re

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document('baocaobtl.docx')

patterns = ['24 ca', '24 use case', '25 collection', '25 bảng', '24 usecase']
for pat in patterns:
    print(f"=== CHECK '{pat}' ===")
    for i, p in enumerate(doc.paragraphs):
        if pat in p.text.lower():
            print(f"  P{i}: {p.text}")
    for t_idx, t in enumerate(doc.tables):
        for r_idx, r in enumerate(t.rows):
            for c_idx, c in enumerate(r.cells):
                if pat in c.text.lower():
                    print(f"  Table {t_idx} R{r_idx} C{c_idx}: {c.text}")
