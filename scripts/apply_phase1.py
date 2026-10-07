import sys
import docx
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl

# Import schema definitions
from collections_data import COLLECTIONS_DATA

sys.stdout.reconfigure(encoding='utf-8')

def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_styled_table(doc, headers, data_rows, col_widths):
    tbl = doc.add_table(rows=len(data_rows) + 1, cols=len(headers), style="Table Grid")
    tbl.autofit = False
    
    # Headers
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "D9D9D9")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.font.bold = True
            
    # Rows
    for r_idx, row_data in enumerate(data_rows):
        row_cells = tbl.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            if c_idx in [1, 2]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.5)
                
    # Column widths
    for row in tbl.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width
            
    return tbl

def main():
    print("Loading baocaobtl.docx...")
    doc = docx.Document("baocaobtl.docx")
    body = doc._body._body

    # ─────────────────────────────────────────────────────────────────────────
    # 1. Sửa Paragraph 32 (Đề cương Chương 3)
    # ─────────────────────────────────────────────────────────────────────────
    for p in doc.paragraphs[:50]:
        if "Chương 3: Thiết kế dữ liệu" in p.text:
            print("Updating paragraph outline for Chapter 3...")
            p.text = "Chương 3: Thiết kế dữ liệu – Thiết kế mô hình CSDL MongoDB gồm 31 collection tiếng Việt chuẩn hóa và bảng hệ thống Users; mô hình quan hệ tham chiếu bằng mã chuỗi nghiệp vụ và cơ chế lưu trữ kép (Dual-Storage); hệ thống chỉ mục (Index) và các ràng buộc toàn vẹn dữ liệu."
            for r in p.runs:
                r.font.name = "Times New Roman"

    # ─────────────────────────────────────────────────────────────────────────
    # 2. Sửa Mục 3.1.2 Kiến trúc CSDL
    # ─────────────────────────────────────────────────────────────────────────
    for p in doc.paragraphs[570:600]:
        if "Cơ sở dữ liệu của hệ thống được đặt tên là" in p.text or "mom_baby_store" in p.text:
            print("Updating 3.1.2 Database architecture...")
            p.text = "Cơ sở dữ liệu của hệ thống được đặt tên là baby_shop_management (theo cấu hình mặc định tại backend/src/config/mongodb.js). Dữ liệu được tổ chức thành 31 collection tiếng Việt chuẩn hóa và bảng hệ thống Users, sử dụng Native MongoDB Driver phiên bản ^7.5.0, hoàn toàn không sử dụng ORM/ODM trung gian như Mongoose nhằm tối ưu tốc độ xử lý dữ liệu và hỗ trợ thao tác trực tiếp với BSON Document."
            for r in p.runs:
                r.font.name = "Times New Roman"

    # ─────────────────────────────────────────────────────────────────────────
    # 3. Thay thế toàn bộ Mục 3.2 (31 Collections)
    # ─────────────────────────────────────────────────────────────────────────
    p_32_heading = None
    p_erd_heading = None
    for child in body:
        if isinstance(child, CT_P):
            p = docx.text.paragraph.Paragraph(child, doc)
            if "3.2. Thiết kế chi tiết các collection" in p.text:
                p_32_heading = child
            elif "3.3. Sơ đồ quan hệ (ERD)" in p.text:
                p_erd_heading = child
                break

    if not p_32_heading or not p_erd_heading:
        raise Exception("Không tìm thấy vị trí bắt đầu Mục 3.2 hoặc Mục 3.3!")

    # Xóa toàn bộ các elements nằm giữa p_32_heading và p_erd_heading
    start_idx = body.index(p_32_heading)
    end_idx = body.index(p_erd_heading)
    print(f"Removing old Section 3.2 elements from index {start_idx + 1} to {end_idx - 1}...")

    elements_to_remove = [body[i] for i in range(start_idx + 1, end_idx)]
    for el in elements_to_remove:
        body.remove(el)

    # Chèn 31 collections mới vào trước p_erd_heading
    print("Inserting 31 new collections with styled tables...")
    headers = ["Tên trường", "Kiểu dữ liệu", "Ràng buộc", "Mô tả"]
    col_widths = [Inches(1.5), Inches(1.1), Inches(1.4), Inches(2.5)]

    for item in COLLECTIONS_DATA:
        # Nếu có nhóm, chèn tiêu đề nhóm
        if item["group"]:
            group_p = doc.add_paragraph()
            group_p.style = "Normal"
            r = group_p.add_run(item["group"])
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            r.font.bold = True
            p_erd_heading.addprevious(group_p._p)

        # Tiêu đề collection
        sub_p = doc.add_paragraph()
        sub_p.style = "Heading 3"
        r = sub_p.add_run(f"{item['num']}. Collection: {item['name']}")
        r.font.name = "Times New Roman"
        p_erd_heading.addprevious(sub_p._p)

        # Mô tả ngắn
        desc_p = doc.add_paragraph()
        desc_p.style = "Normal"
        r = desc_p.add_run(item["desc"])
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        r.font.italic = True
        p_erd_heading.addprevious(desc_p._p)

        # Bảng dữ liệu
        tbl = create_styled_table(doc, headers, item["fields"], col_widths)
        p_erd_heading.addprevious(tbl._tbl)

        # Đoạn trắng cách bảng
        sep_p = doc.add_paragraph()
        p_erd_heading.addprevious(sep_p._p)

    # ─────────────────────────────────────────────────────────────────────────
    # 4. Cập nhật Mục 3.3 (Sơ đồ quan hệ ERD)
    # ─────────────────────────────────────────────────────────────────────────
    print("Updating 3.3 ERD text...")
    for p in doc.paragraphs:
        if "Hệ thống được thiết kế với các mối quan hệ (relationships) chủ yếu như sau:" in p.text:
            p.text = "Hệ thống được thiết kế với mô hình cơ sở dữ liệu phi quan hệ (Document Store) linh hoạt, trong đó các thực thể nghiệp vụ liên kết chặt chẽ thông qua các quy ước kiến trúc:"
        elif "Nhân viên và Vai trò (1 - N):" in p.text:
            p.text = "1. Liên kết bằng Mã chuỗi nghiệp vụ (String Business Code): Hệ thống sử dụng các mã định danh tự sinh qua bộ đếm Counters (MaSP, MaKH, MaNCC, MaHD, MaPN, MaDDH...) làm khóa liên kết logic giữa các collection thay vì phụ thuộc cứng nhắc vào ObjectId."
        elif "Sản phẩm và Loại hàng/Danh mục (1 - N):" in p.text:
            p.text = "2. Cơ chế lưu trữ kép chi tiết chứng từ (Dual-Storage Pattern): Danh mục hàng hóa trong chứng từ vừa được nhúng trực tiếp (embedded) trong document cha phục vụ hiển thị và in ấn nhanh, vừa được tách ra lưu tại 8 collection chi tiết CT_* phục vụ phân tích báo cáo và kiểm kê."
        elif "Đơn đặt hàng và Nhà cung cấp, Nhân viên (1 - N):" in p.text:
            p.text = "3. Nhân viên và Vai trò (N - 1): Mỗi nhân viên liên kết với một vai trò qua MaVaiTro (1-5), phân quyền chặt chẽ theo ma trận RBAC."
        elif "Đơn hàng (Bán) và Khách hàng, Nhân viên (1 - N):" in p.text:
            p.text = "4. Sản phẩm và Loại hàng (N - 1): Mỗi sản phẩm thuộc về một loại hàng (MaLoai), đảm bảo tính phân nhóm kế toán."
        elif "Chứng từ và Chi tiết (1 - N): Giữa Hóa đơn và CT_Hóa đơn" in p.text:
            p.text = "5. Luồng Mua hàng - Nhập kho - Công nợ NCC: Đơn đặt hàng (DonDatHang) liên kết Nhà cung cấp -> Phiếu nhập kho (PhieuNhap) kế thừa đơn mua, tự động tăng TonKho và phát sinh sổ nợ (CongNo) -> Phiếu chi (PhieuChi) thanh toán công nợ."
        elif "Tồn kho và Sản phẩm (1 - 1): Quản lý tồn kho theo từng mã sản phẩm." in p.text:
            p.text = "6. Luồng Bán hàng - Xuất kho - Công nợ Khách hàng: Đơn hàng POS (DonHang) liên kết Khách hàng -> Hóa đơn (HoaDon) kiêm chứng từ thu tiền -> Phiếu xuất kho (PhieuXuat) xuất hàng -> Phiếu thu (PhieuThu) ghi nhận thu quỹ tiền mặt."

    # ─────────────────────────────────────────────────────────────────────────
    # 5. Cập nhật Mục 3.4 (Ràng buộc dữ liệu và Validation)
    # ─────────────────────────────────────────────────────────────────────────
    print("Updating 3.4 Validation rules...")
    for p in doc.paragraphs:
        if "Trong MongoDB, việc kiểm tra tính toàn vẹn và hợp lệ của dữ liệu" in p.text:
            p.text = "Trong hệ thống Baby Shop Management ERP, cơ chế kiểm tra tính toàn vẹn và hợp lệ của dữ liệu được kết hợp chặt chẽ giữa 2 tầng: Tầng chỉ mục duy nhất (Database Indexes) và Tầng kiểm soát nghiệp vụ ứng dụng (Controller / Middleware Validation):"
        elif "Khóa chính tự động: Tất cả các bảng sử dụng kiểu ObjectId" in p.text:
            p.text = "• Khóa chính và Khóa nghiệp vụ: Tất cả các collection sử dụng ObjectId làm khóa chính tự sinh; đồng thời các mã nghiệp vụ (MaSP, MaHD, MaPN, MaKH...) được bảo đảm duy nhất thông qua cơ chế atomic $inc của collection Counters."
        elif "Khóa ngoại: Sử dụng tham chiếu ObjectId sang các bảng liên quan." in p.text:
            p.text = "• Toàn vẹn tham chiếu và Chặn xóa dữ liệu ràng buộc: Hệ thống kiểm tra ràng buộc trước khi xóa: Không cho phép xóa Loại hàng nếu đang có sản phẩm thuộc danh mục; đối với Nhà cung cấp, Khách hàng và Sản phẩm đã phát sinh chứng từ, hệ thống tự động chuyển sang cơ chế xóa mềm (Soft-delete: 'Ngưng hoạt động' hoặc 'Ngừng bán')."
        elif "Bắt buộc (Required): Các trường quan trọng như Tên, Mã" in p.text:
            p.text = "• Chặn xóa vật lý chứng từ kế toán: Các chứng từ gốc (PhieuNhap, HoaDon, DonHang, PhieuXuat, DonDatHang) được bảo vệ bằng danh sách PROTECTED_VOUCHERS, tuyệt đối không được xóa vật lý mà chỉ được hủy bỏ qua trạng thái chứng từ."
        elif "Unique (Duy nhất): Các trường như username, mã khách hàng" in p.text:
            p.text = "• Ràng buộc định dạng và nghiệp vụ bắt buộc: Đơn giá và số lượng phải là số dương (>0); sản phẩm bắt buộc có Hạn sử dụng (HanSuDung); CCCD phải gồm đúng 12 chữ số; số điện thoại phải gồm 10-11 chữ số và bắt đầu bằng số 0."
        elif "Kiểu dữ liệu: Được kiểm soát nghiêm ngặt bằng String, Number" in p.text:
            p.text = "• Cơ chế chặn bán âm kho và Rollback giao dịch: Hàm adjustStock kiểm tra điều kiện tồn kho tức thời ($gte: quantity). Nếu tồn không đủ, hệ thống lập tức ném lỗi 409 Conflict và thực hiện rollback hoàn trả lại các mặt hàng đã trừ trong giao dịch, bảo đảm an toàn dữ liệu tuyệt đối kể cả trên MongoDB Standalone."

    # ─────────────────────────────────────────────────────────────────────────
    # 6. Cập nhật Mục 3.6 (Chỉ mục Indexes)
    # ─────────────────────────────────────────────────────────────────────────
    print("Updating 3.6 Indexes...")
    for p in doc.paragraphs:
        if "Hóa đơn: { createdAt: -1 }, { MaKH: 1, createdAt: -1 }" in p.text:
            p.text = "• HoaDon: { createdAt: -1 }, { MaKH: 1, createdAt: -1 }, { TrangThai: 1, createdAt: -1 }, { NgayLap: -1 }, { MaHD: 1 } (unique, sparse)"
        elif "Đơn hàng: { createdAt: -1 }, { MaKH: 1, createdAt: -1 }" in p.text:
            p.text = "• DonHang: { createdAt: -1 }, { MaKH: 1, createdAt: -1 }, { TrangThai: 1, createdAt: -1 }, { NgayDat: -1 }, { MaDH: 1 } (unique, sparse)"
        elif "Phiếu nhập: { createdAt: -1 }, { MaNCC: 1, createdAt: -1 }" in p.text:
            p.text = "• PhieuNhap: { createdAt: -1 }, { MaNCC: 1, createdAt: -1 }, { NgayNhap: -1 }, { MaPN: 1 } (unique, sparse)\n• DonDatHang: { MaNCC: 1, createdAt: -1 }, { TrangThai: 1, createdAt: -1 }\n• PhieuXuat: { createdAt: -1 }, { NgayXuat: -1 }"
        elif "Công nợ: { MaPN: 1 }, { MaNCC: 1, TrangThai: 1 }" in p.text:
            p.text = "• CongNo: { MaPN: 1 } (lookup cascade), { MaNCC: 1, TrangThai: 1 }, { TrangThai: 1, createdAt: -1 }\n• PhieuThu & PhieuChi: { createdAt: -1 }"
        elif "Tồn kho: { MaSP: 1 } (unique)" in p.text:
            p.text = "• TonKho: { MaSP: 1 } (unique, sparse, quan hệ 1-1 với SanPham)"
        elif "Sản phẩm: { MaSP: 1 }, { TrangThai: 1 }, { MaLoai: 1 }" in p.text:
            p.text = "• SanPham: { MaSP: 1 } (unique, sparse), { TrangThai: 1 }, { MaLoai: 1 }, { TenSP: 'text' } (Full-text search tên sản phẩm)"
        elif "Khách hàng: { MaKH: 1 }, { SDT: 1 }" in p.text:
            p.text = "• KhachHang: { MaKH: 1 } (unique, sparse), { SDT: 1 } (unique, sparse), { TrangThai: 1 }\n• NhaCungCap: { MaNCC: 1 } (unique, sparse), { TrangThai: 1 }"
        elif "Nhà cung cấp: { MaNCC: 1 }, { TrangThai: 1 }" in p.text:
            p.text = "• Chi tiết chứng từ (CT_*): CT_HoaDon { MaHD: 1 }, CT_DonHang { MaDH: 1 }, CT_PhieuNhap { MaPN: 1 }, CT_PhieuXuat { MaPX: 1 }, CT_DonDatHang { MaDDH: 1 }"
        elif "Nhật ký kiểm toán (Audit Logs): { timestamp: -1 }" in p.text:
            p.text = "• Users: { username: 1 } (unique, sparse)\n• AuditLogs: { timestamp: -1 }, { action: 1, timestamp: -1 }, { module: 1, timestamp: -1 }, { role: 1, timestamp: -1 }\n• Counters: { _id: 1 }"

    # ─────────────────────────────────────────────────────────────────────────
    # 7. Cập nhật Mục 3.7 (Script tạo CSDL)
    # ─────────────────────────────────────────────────────────────────────────
    print("Updating 3.7 Database script text...")
    for p in doc.paragraphs:
        if "Hệ thống sử dụng script SQL" in p.text:
            p.text = "Hệ thống sử dụng mã nguồn JavaScript tự động khởi tạo cơ sở dữ liệu và chỉ mục bằng Native MongoDB Driver (thông qua tệp ensure-indexes.js và seed.js khi khởi động máy chủ). Dưới đây là đoạn mã khởi tạo tiêu biểu:"
        elif 'db.createCollection("users"' in p.text:
            p.text = 'await db.collection("Users").createIndex({ username: 1 }, { unique: true, sparse: true });\nawait db.collection("SanPham").createIndex({ MaSP: 1 }, { unique: true, sparse: true });\nawait db.collection("HoaDon").createIndex({ MaHD: 1 }, { unique: true, sparse: true });\nawait db.collection("AuditLogs").createIndex({ timestamp: -1 });'

    # ─────────────────────────────────────────────────────────────────────────
    # 8. Cập nhật Table 34 (Tổng hợp các module và bảng)
    # ─────────────────────────────────────────────────────────────────────────
    print("Updating Table 34 (summary of modules and tables)...")
    for t in doc.tables:
        if len(t.rows) > 0 and len(t.rows[0].cells) >= 4 and "Bảng dữ liệu chính" in t.rows[0].cells[3].text:
            rows = t.rows
            if len(rows) >= 9:
                rows[1].cells[3].text = "Users, VaiTro"
                rows[2].cells[3].text = "KhachHang, NhaCungCap, SanPham, LoaiHang, NhanVien"
                rows[3].cells[3].text = "DonDatHang, CT_DonDatHang, PhieuNhap, CT_PhieuNhap, TonKho"
                rows[4].cells[3].text = "DonHang, CT_DonHang, HoaDon, CT_HoaDon, PhieuXuat, CT_PhieuXuat, TonKho, KhachHang"
                rows[5].cells[3].text = "PhieuThu, PhieuChi, ThanhToan, CongNo"
                rows[6].cells[3].text = "TonKho, SanPham, KiemKe, CT_KiemKe, DieuChinhKho, PhieuTraHang, CT_PhieuTraHang"
                rows[7].cells[3].text = "Tổng hợp động từ các bảng chứng từ gốc trên"
                rows[8].cells[3].text = "AuditLogs, Counters, KhuyenMai, CT_KhuyenMai, Tệp sao lưu JSON (backups/)"
                rows[8].cells[4].text = "Tìm kiếm nhanh; Lọc dữ liệu; Phân trang; In hóa đơn theo mẫu kế toán; Xuất Excel (.xls XML)/CSV; Thống kê biểu đồ SVG; Thanh toán VietQR động chuẩn NAPAS; Gửi email hóa đơn; Nhật ký kiểm toán; Sao lưu và phục hồi dữ liệu."

    # ─────────────────────────────────────────────────────────────────────────
    # 9. Cập nhật Use Cases (UC01 - UC24)
    # ─────────────────────────────────────────────────────────────────────────
    print("Updating Use Cases UC01 - UC24 technologies and API routes...")
    for p in doc.paragraphs[145:550]:
        text = p.text

        # bcrypt -> scrypt
        if "bcrypt" in text:
            p.text = text.replace("bcryptjs", "scrypt (Node.js crypto)").replace("bcrypt", "scrypt")
            text = p.text

        # Recharts -> Custom SVG Charts
        if "Recharts" in text or "recharts" in text:
            p.text = text.replace("thư viện Recharts", "các thành phần biểu đồ SVG thuần React (LineChart, BarChart, DonutChart trong Charts.jsx)").replace("Recharts", "Custom SVG Charts")
            text = p.text

        # xlsx -> exportToExcel (.xls XML)
        if "xlsx" in text:
            p.text = text.replace("thông qua thư viện xlsx", "thông qua tiện ích xuất bảng tính HTML/XML Spreadsheet (.xls) và CSV (exportToExcel, exportToCsv) kết hợp thư viện jsPDF và phương thức in ấn trực tiếp window.print()").replace("xlsx", "exportToExcel (.xls XML)")
            text = p.text

        # Đường dẫn API trong các Use Case
        if "POST /api/accounts" in text:
            p.text = text.replace("POST /api/accounts", "POST /api/admin/accounts")
        elif "GET /api/accounts" in text:
            p.text = text.replace("GET /api/accounts", "GET /api/admin/accounts")
        elif "PUT /api/accounts/:id" in text:
            p.text = text.replace("PUT /api/accounts/:id", "PUT /api/admin/accounts/:id (và PATCH toggle-lock)")
        elif "POST /api/categories" in text:
            p.text = text.replace("POST /api/categories", "POST /api/product-categories")
        elif "GET /api/categories" in text:
            p.text = text.replace("GET /api/categories", "GET /api/product-categories")
        elif "POST /api/employees" in text:
            p.text = text.replace("POST /api/employees", "POST /api/admin/employees")
        elif "GET /api/employees" in text:
            p.text = text.replace("GET /api/employees", "GET /api/admin/employees")
        elif "POST /api/warehouse-receipts" in text:
            p.text = text.replace("POST /api/warehouse-receipts", "POST /api/goods-receipts (lập phiếu nhập kho, tăng tồn kho và ghi nợ NCC)")
        elif "POST /api/sales-invoices" in text:
            p.text = text.replace("POST /api/sales-invoices", "POST /api/sales-orders (tạo đơn bán lẻ POS, trừ tồn kho, tự sinh Hóa đơn HoaDon và Phiếu xuất PhieuXuat)")
        elif "GET /api/sales-invoices" in text:
            p.text = text.replace("GET /api/sales-invoices", "GET /api/invoices (tra cứu danh sách hóa đơn)")
        elif "POST /api/warehouse-exports" in text:
            p.text = text.replace("POST /api/warehouse-exports", "POST /api/goods-issues")
        elif "GET /api/warehouse-exports" in text:
            p.text = text.replace("GET /api/warehouse-exports", "GET /api/goods-issues")
        elif "POST /api/stock-takes/:id/complete" in text:
            p.text = text.replace("POST /api/stock-takes/:id/complete", "POST /api/stocktakes/:id/adjust (tự động cân bằng tồn kho thực tế)")
        elif "GET /api/inventory/alerts" in text:
            p.text = "2. Phía frontend tự động lọc từ danh sách tồn kho (GET /api/inventory) các sản phẩm có số lượng chạm ngưỡng an toàn (stock <= SoLuongToiThieu hoặc LOW_STOCK_THRESHOLD = 10) và quét sản phẩm cận hạn sử dụng (<60 ngày) qua thanh thông báo NotificationPopover."
        elif "POST /api/products/upload-image" in text:
            p.text = "4. Người dùng tải ảnh sản phẩm từ máy tính: Trình duyệt nén ảnh qua HTML5 Canvas thành chuỗi Base64 Data URL và gửi trực tiếp trong payload POST /api/products lưu vào trường HinhAnh của SanPham."

    # ─────────────────────────────────────────────────────────────────────────
    # 10. Cập nhật Table 8 (Danh mục 24 Use Case)
    # ─────────────────────────────────────────────────────────────────────────
    print("Updating Table 8 (Use Case table)...")
    for t in doc.tables:
        if len(t.rows) > 0 and len(t.rows[0].cells) >= 5 and "Mã UC" in t.rows[0].cells[0].text:
            for r in t.rows[1:]:
                uc_code = r.cells[0].text.strip()
                if uc_code == "UC04":
                    r.cells[4].text = "AdminPage (tab Người dùng)"
                elif uc_code == "UC05":
                    r.cells[4].text = "CustomersPage (Khách hàng & Tích điểm)"
                elif uc_code == "UC06":
                    r.cells[4].text = "SuppliersPage (Nhà cung cấp)"
                elif uc_code == "UC07":
                    r.cells[4].text = "ProductsPage (Sản phẩm & HSD)"
                elif uc_code == "UC08":
                    r.cells[4].text = "CategoriesPage (Loại hàng)"
                elif uc_code == "UC09":
                    r.cells[4].text = "AdminPage (tab Nhân viên)"
                elif uc_code == "UC10" or uc_code == "UC12":
                    r.cells[4].text = "PurchaseOrderPage (Đặt hàng NCC)"
                elif uc_code == "UC11":
                    r.cells[4].text = "WarehouseDocumentsPage (tab Nhập kho)"
                elif uc_code == "UC13":
                    r.cells[4].text = "SalesPOSPage (Bán hàng POS tại quầy)"
                elif uc_code == "UC14":
                    r.cells[4].text = "InvoicePage (Hóa đơn & Thanh toán)"
                elif uc_code == "UC15" or uc_code == "UC16":
                    r.cells[4].text = "WarehouseDocumentsPage (tab Xuất kho)"
                elif uc_code == "UC17":
                    r.cells[4].text = "CashVouchersPage (tab Phiếu thu 01-TT)"
                elif uc_code == "UC18":
                    r.cells[4].text = "CashVouchersPage (tab Phiếu chi 02-TT)"
                elif uc_code == "UC19":
                    r.cells[4].text = "CashVouchersPage & DebtsPage"
                elif uc_code == "UC20":
                    r.cells[4].text = "InventoryPage (Tồn kho thời gian thực)"
                elif uc_code == "UC21":
                    r.cells[4].text = "StocktakePage (Kiểm kê & Cân bằng kho)"
                elif uc_code == "UC22":
                    r.cells[4].text = "InventoryPage & NotificationPopover"
                elif uc_code == "UC23":
                    r.cells[4].text = "DashboardPage (Tổng quan KPI)"
                elif uc_code == "UC24":
                    r.cells[4].text = "ReportPage (Hệ thống 17 Báo cáo)"

    # Lưu kết quả ra baocaobtl_fixed.docx
    output_filename = "baocaobtl_fixed.docx"
    print(f"Saving changes to {output_filename}...")
    doc.save(output_filename)
    print("SUCCESS: Phase 1 completed successfully!")

if __name__ == "__main__":
    main()
